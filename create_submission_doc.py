import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_document():
    doc = docx.Document()

    # Configure Margins for 2-page fit
    for section in doc.sections:
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # Styling helpers
    navy = RGBColor(26, 54, 138)
    dark_gray = RGBColor(51, 65, 85)
    charcoal = RGBColor(15, 23, 42)

    def set_font(run, name="Calibri", size=9.5, bold=False, color=charcoal):
        run.font.name = name
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color

    def add_para(doc, text="", space_after=2, line_spacing=1.05):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            r = p.add_run(text)
            set_font(r)
        return p

    def set_cell_background(cell, fill_hex):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
            node = OxmlElement(m)
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    # 1. Header Table (Student details)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    cell_l = table.rows[0].cells[0]
    cell_r = table.rows[0].cells[1]
    cell_l.width = Inches(4.5)
    cell_r.width = Inches(3.0)

    set_cell_background(cell_l, "F1F5F9")
    set_cell_background(cell_r, "F1F5F9")
    set_cell_margins(cell_l, top=70, bottom=70, left=120, right=120)
    set_cell_margins(cell_r, top=70, bottom=70, left=120, right=120)

    p_l = cell_l.paragraphs[0]
    p_l.paragraph_format.space_after = Pt(0)
    r1 = p_l.add_run("Name: ")
    set_font(r1, size=9, bold=True)
    r2 = p_l.add_run("Aakash Valliyil   ")
    set_font(r2, size=9)
    r3 = p_l.add_run("Roll No: ")
    set_font(r3, size=9, bold=True)
    r4 = p_l.add_run("763   ")
    set_font(r4, size=9)
    r5 = p_l.add_run("Division: ")
    set_font(r5, size=9, bold=True)
    r6 = p_l.add_run("A")
    set_font(r6, size=9)

    p_r = cell_r.paragraphs[0]
    p_r.paragraph_format.space_after = Pt(0)
    p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r7 = p_r.add_run("DOP: ")
    set_font(r7, size=9, bold=True)
    r8 = p_r.add_run("03/10/2026   ")
    set_font(r8, size=9)
    r9 = p_r.add_run("DOS: ")
    set_font(r9, size=9, bold=True)
    r10 = p_r.add_run("06/10/2026   ")
    set_font(r10, size=9)
    r11 = p_r.add_run("Score: _____")
    set_font(r11, size=9, bold=True)

    # 2. Title & Metadata
    p_title = add_para(doc, space_after=1)
    r_title = p_title.add_run("Project Title: MapPulse AI – Google Maps Competitor Update Intelligence Platform")
    set_font(r_title, size=11.5, bold=True, color=navy)

    p_meta = add_para(doc, space_after=3)
    p_meta.add_run("• Live Application URL: ").font.bold = True
    p_meta.add_run("https://google-maps-competitor-intelligence-c7ga.onrender.com\n")
    p_meta.add_run("• GitHub Repository: ").font.bold = True
    p_meta.add_run("https://github.com/AVR003/google-maps-competitor-intelligence   ")
    p_meta.add_run("• Demo Subject: ").font.bold = True
    p_meta.add_run("ABC Salon, Kharghar (vs 4 Competitors)")
    for run in p_meta.runs:
        if not run.bold:
            set_font(run, size=8.5, color=dark_gray)
        else:
            set_font(run, size=8.5, bold=True, color=navy)

    # 3. Approach & Methodology
    p_h1 = add_para(doc, space_after=1)
    set_font(p_h1.add_run("1. Approach & Methodology"), size=10, bold=True, color=navy)

    p_app = add_para(doc, space_after=2)
    set_font(p_app.add_run(
        "MapPulse AI is an automated intelligence tool engineered to monitor competitor Google Maps updates, analyze promotional topics, detect gaps, and generate non-repetitive, ready-to-publish Google Maps posts."
    ), size=8.5)

    workflow_items = [
        ("Multi-Project Setup: ", "Groups local business profiles, competitors, and keywords. Pre-loaded with a realistic demo dataset for 100% evaluation reliability."),
        ("Automated Scraping & Anti-Automation: ", "Headless Python Selenium engine with stealth headers extracts post copy, timestamps, images, and CTAs. If Google presents a CAPTCHA challenge, the system halts gracefully, prompts the user via a modal, and logs manual intervention."),
        ("SHA-256 Duplicate Detection: ", "Fingerprints posts via SHA256(competitor + text + date). Re-scraping identifies and skips existing posts to maintain clean data."),
        ("AI Trend & Gap Analysis: ", "Computes the Topic Occurrence Matrix (% occurrence, competitor usage count), analyzes publishing frequency, and identifies untapped market gaps."),
        ("Content & Update Generator: ", "Allows users to control idea batch size (3, 10, 50). Employs duplicate prevention against SQLite history and generates complete copy with location keywords, CTAs, and AI image concepts.")
    ]

    for idx, (title, desc) in enumerate(workflow_items, start=1):
        pw = add_para(doc, space_after=1.5)
        rt = pw.add_run(f"{idx}. {title}")
        set_font(rt, size=8.5, bold=True)
        rd = pw.add_run(desc)
        set_font(rd, size=8.5)

    # 4. Technologies Used
    p_h2 = add_para(doc, space_after=1)
    set_font(p_h2.add_run("2. Technologies Used"), size=10, bold=True, color=navy)

    p_tech = add_para(doc, space_after=3)
    tech_runs = [
        ("Backend & DB: ", "Python 3.14, FastAPI, Uvicorn, SQLite.   "),
        ("Scraping: ", "Selenium WebDriver, BeautifulSoup4.   \n"),
        ("Frontend: ", "HTML5, Tailwind CSS, Chart.js, Lucide Icons.   "),
        ("AI Layer: ", "Google Gemini 2.5 Flash, Groq LLaMA 3.3, Built-in Offline NLP.   "),
        ("Hosting: ", "Render.com (Procfile).")
    ]
    for lbl, val in tech_runs:
        r_l = p_tech.add_run(lbl)
        set_font(r_l, size=8.5, bold=True)
        r_v = p_tech.add_run(val)
        set_font(r_v, size=8.5)

    # 5. Screenshots Part 1 (2-column layout on Page 1)
    p_h3_1 = add_para(doc, space_after=1)
    set_font(p_h3_1.add_run("3. Key Screenshots (Part 1: Ingestion & Duplicate Auditing)"), size=10, bold=True, color=navy)

    img_dir = r"C:\Users\aakas\Desktop\google-maps-competitor-intelligence\extracted_images\word\media"
    img1_path = os.path.join(img_dir, "image1.png")
    img2_path = os.path.join(img_dir, "image2.png")
    img3_path = os.path.join(img_dir, "image3.png")
    img4_path = os.path.join(img_dir, "image4.png")

    tbl_ss1 = doc.add_table(rows=1, cols=2)
    tbl_ss1.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ss1.autofit = False

    c1_1 = tbl_ss1.rows[0].cells[0]
    c1_2 = tbl_ss1.rows[0].cells[1]
    c1_1.width = Inches(3.6)
    c1_2.width = Inches(3.6)

    # Cell 1: Dashboard
    p1 = c1_1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(1)
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c1 = p1.add_run("Fig 1: Dashboard & Analytics Overview\n")
    set_font(r_c1, size=8, bold=True)
    if os.path.exists(img1_path):
        p1.add_run().add_picture(img1_path, width=Inches(3.5))
    p1_sub = c1_1.add_paragraph()
    p1_sub.paragraph_format.space_after = Pt(0)
    p1_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p1_sub.add_run("Monitors 4 competitors, 29 stored posts, and 13 duplicates skipped."), size=7.5, color=dark_gray)

    # Cell 2: Competitors & Scraper
    p2 = c1_2.paragraphs[0]
    p2.paragraph_format.space_after = Pt(1)
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c2 = p2.add_run("Fig 2: Competitor Management & Scraper Console\n")
    set_font(r_c2, size=8, bold=True)
    if os.path.exists(img2_path):
        p2.add_run().add_picture(img2_path, width=Inches(3.5))
    p2_sub = c1_2.add_paragraph()
    p2_sub.paragraph_format.space_after = Pt(0)
    p2_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p2_sub.add_run("Live Selenium console with SHA-256 duplicate detection."), size=7.5, color=dark_gray)

    # PAGE BREAK TO EXACT PAGE 2
    doc.add_page_break()

    # ---------------- PAGE 2 ----------------
    p_h3_2 = add_para(doc, space_after=1)
    set_font(p_h3_2.add_run("3. Key Screenshots (Part 2: Trend Intelligence & Content Generation)"), size=10, bold=True, color=navy)

    tbl_ss2 = doc.add_table(rows=1, cols=2)
    tbl_ss2.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ss2.autofit = False

    c2_1 = tbl_ss2.rows[0].cells[0]
    c2_2 = tbl_ss2.rows[0].cells[1]
    c2_1.width = Inches(3.6)
    c2_2.width = Inches(3.6)

    # Cell 3: Trend Analysis
    p3 = c2_1.paragraphs[0]
    p3.paragraph_format.space_after = Pt(1)
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c3 = p3.add_run("Fig 3: AI Trend Matrix & Strategic Gap Analysis\n")
    set_font(r_c3, size=8, bold=True)
    if os.path.exists(img3_path):
        p3.add_run().add_picture(img3_path, width=Inches(3.5))
    p3_sub = c2_1.add_paragraph()
    p3_sub.paragraph_format.space_after = Pt(0)
    p3_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p3_sub.add_run("Assignment Table 13 matrix (Festival Offer 33%, Hair Care Tips 28%)."), size=7.5, color=dark_gray)

    # Cell 4: Generator
    p4 = c2_2.paragraphs[0]
    p4.paragraph_format.space_after = Pt(1)
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c4 = p4.add_run("Fig 4: Content Generator & Ready-to-Publish Update\n")
    set_font(r_c4, size=8, bold=True)
    if os.path.exists(img4_path):
        p4.add_run().add_picture(img4_path, width=Inches(3.5))
    p4_sub = c2_2.add_paragraph()
    p4_sub.paragraph_format.space_after = Pt(0)
    p4_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p4_sub.add_run("Controlled batch size generator drafting full copy and photo concepts."), size=7.5, color=dark_gray)

    # 6. Key Results & Findings
    p_h4 = add_para(doc, space_after=1)
    p_h4.paragraph_format.space_before = Pt(4)
    set_font(p_h4.add_run("4. Key Results & Findings from Competitor Research"), size=10, bold=True, color=navy)

    findings = [
        ("Topic Saturation: ", "Competitors heavily emphasize Festival Offers (33.3%) and Hair Care Tips (27.8%), reflecting strong local demand for discount bundles and seasonal hair damage prevention."),
        ("Identified Market Gap: ", "Visual Before/After transformations (22.2%) and specialized Bridal Packages (5.6%) are underutilized by 50% of competitors, presenting a primary opportunity for ABC Salon to capture local search visibility."),
        ("Publishing Patterns: ", "Competitors post primarily mid-week (Tuesday–Thursday) leading into weekend walk-in traffic, with CTA distributions concentrated on 'Book Online' (34%) and 'Call Now' (28%).")
    ]
    for lbl, desc in findings:
        pf = add_para(doc, space_after=1.5)
        rf_l = pf.add_run(f"• {lbl}")
        set_font(rf_l, size=8.5, bold=True)
        rf_d = pf.add_run(desc)
        set_font(rf_d, size=8.5)

    # 7. Bonus Features Implemented
    p_h5 = add_para(doc, space_after=1)
    p_h5.paragraph_format.space_before = Pt(4)
    set_font(p_h5.add_run("5. Bonus Features Implemented"), size=10, bold=True, color=navy)

    bonuses = [
        ("Keyword-Based Competitor Discovery: ", "Discovers candidate businesses from search keywords (e.g., 'salon in Kharghar') and adds them directly to the scraping queue without duplicates."),
        ("Strategic Gap Analysis & Recommendations: ", "Automatically prescribes content strategies based on what competitors have missed."),
        ("Offline Resilience: ", "The integrated heuristic NLP engine guarantees full functionality even without external AI API keys or internet access.")
    ]
    for lbl, desc in bonuses:
        pb = add_para(doc, space_after=1.5)
        rb_l = pb.add_run(f"• {lbl}")
        set_font(rb_l, size=8.5, bold=True)
        rb_d = pb.add_run(desc)
        set_font(rb_d, size=8.5)

    output_path = r"C:\Users\aakas\Desktop\Aakash_Valliyil_763_Google_Maps_Intelligence_Submission.docx"
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    create_document()
